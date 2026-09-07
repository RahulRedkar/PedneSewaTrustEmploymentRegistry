"""
Data Import Utility for Pedne Sewa Trust.
Supports importing Visiting Register and candidate datasets from Excel (.xlsx) and CSV files.
Provides intelligent fuzzy header mapping and robust date/time formatting.
"""

import csv
import os
from datetime import datetime, date, time
from pathlib import Path
from typing import List, Dict, Any, Optional

from database.repository import repository
from utils.logger import logger
from utils.validators import clean_mobile


HEADER_SYNONYMS = {
    "sr_no": ["sr no", "sr. no.", "sr no.", "serial no", "serial number", "s.no", "s no", "sno", "id"],
    "visit_date": ["date", "visit date", "date of visit", "visiting date", "entry date"],
    "visit_time": ["time", "visit time", "time of visit", "in time", "visiting time"],
    "candidate_name": ["name of candidate", "candidate name", "name", "visitor name", "full name", "candidate"],
    "village": ["address (village)", "address(village)", "address", "village", "location", "residence", "area"],
    "mobile": ["mobile no.", "mobile no", "mobile", "mobile number", "contact no", "contact", "phone", "phone no"],
    "purpose": ["purpose of visit", "purpose", "reason", "visit purpose", "work"],
    "remarks": ["remarks", "remark", "notes", "comment", "comments", "other remarks", "details"]
}


def _match_header_key(raw_header: str) -> Optional[str]:
    """Normalizes header and matches against known field synonyms."""
    cleaned = str(raw_header or "").strip().lower().replace("_", " ")
    # Exact match check
    for canonical_key, synonyms in HEADER_SYNONYMS.items():
        if cleaned == canonical_key:
            return canonical_key
        for syn in synonyms:
            if cleaned == syn:
                return canonical_key

    # Substring check
    for canonical_key, synonyms in HEADER_SYNONYMS.items():
        for syn in synonyms:
            if syn in cleaned:
                return canonical_key
    return None


def _format_cell_date(val: Any) -> str:
    """Converts various date representations into YYYY-MM-DD string."""
    if val is None:
        return datetime.now().strftime("%Y-%m-%d")
    if isinstance(val, (datetime, date)):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return datetime.now().strftime("%Y-%m-%d")
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return s


def _format_cell_time(val: Any) -> str:
    """Converts various time representations into hh:mm AM/PM string."""
    if val is None:
        return datetime.now().strftime("%I:%M %p")
    if isinstance(val, datetime):
        return val.strftime("%I:%M %p")
    if isinstance(val, time):
        return val.strftime("%I:%M %p")
    s = str(val).strip()
    if not s:
        return datetime.now().strftime("%I:%M %p")
    for fmt in ("%H:%M:%S", "%H:%M", "%I:%M %p", "%I:%M:%S %p"):
        try:
            return datetime.strptime(s, fmt).strftime("%I:%M %p")
        except ValueError:
            pass
    return s


class FileImporter:
    """Parses and imports records from Excel and CSV files into the database."""

    def __init__(self, repo=None):
        self.repo = repo or repository

    @staticmethod
    def parse_visiting_register_file(filepath: str) -> List[Dict[str, Any]]:
        """
        Reads an Excel or CSV file and extracts Visiting Register records.
        Returns a list of parsed record dicts.
        """
        path = Path(filepath)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {filepath}")

        ext = path.suffix.lower()
        if ext in (".xlsx", ".xlsm", ".xltx"):
            return FileImporter._parse_excel(path)
        elif ext in (".csv", ".txt"):
            return FileImporter._parse_csv(path)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Please provide an Excel (.xlsx) or CSV (.csv) file.")

    @staticmethod
    def _parse_excel(path: Path) -> List[Dict[str, Any]]:
        """Parses Excel file using openpyxl."""
        import openpyxl
        wb = openpyxl.load_workbook(str(path), data_only=True)

        # Select 'Visiting Register' sheet if available, otherwise active sheet
        sheet = None
        for name in wb.sheetnames:
            if "visit" in name.lower():
                sheet = wb[name]
                break
        if not sheet:
            sheet = wb.active

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        # Header row detection: find first row with at least 2 non-empty cells
        header_row_idx = -1
        col_map: Dict[int, str] = {}
        for idx, r in enumerate(rows):
            non_empty = [c for c in r if c is not None and str(c).strip() != ""]
            if len(non_empty) >= 2:
                # Check for matches
                matches = 0
                temp_map = {}
                for col_idx, cell in enumerate(r):
                    matched = _match_header_key(cell)
                    if matched:
                        temp_map[col_idx] = matched
                        matches += 1
                if matches >= 2:
                    header_row_idx = idx
                    col_map = temp_map
                    break

        if header_row_idx == -1:
            # Fallback: treat Row 0 as headers
            header_row_idx = 0
            for col_idx, cell in enumerate(rows[0]):
                matched = _match_header_key(cell)
                if matched:
                    col_map[col_idx] = matched

        records = []
        for r in rows[header_row_idx + 1:]:
            if not any(c is not None and str(c).strip() != "" for c in r):
                continue  # Skip completely blank rows

            item: Dict[str, Any] = {}
            for col_idx, key in col_map.items():
                if col_idx < len(r):
                    val = r[col_idx]
                    if key == "visit_date":
                        item[key] = _format_cell_date(val)
                    elif key == "visit_time":
                        item[key] = _format_cell_time(val)
                    elif key == "mobile":
                        item[key] = clean_mobile(str(val or ""))
                    else:
                        item[key] = str(val or "").strip()

            if item.get("candidate_name"):
                records.append(item)

        return records

    @staticmethod
    def _parse_csv(path: Path) -> List[Dict[str, Any]]:
        """Parses CSV file with UTF-8 BOM fallback."""
        content = ""
        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                with open(str(path), "r", encoding=encoding) as f:
                    content = f.read()
                break
            except UnicodeDecodeError:
                continue

        reader = list(csv.reader(content.splitlines()))
        if not reader:
            return []

        header_row_idx = -1
        col_map: Dict[int, str] = {}
        for idx, r in enumerate(reader):
            non_empty = [c for c in r if c and c.strip()]
            if len(non_empty) >= 2:
                matches = 0
                temp_map = {}
                for col_idx, cell in enumerate(r):
                    matched = _match_header_key(cell)
                    if matched:
                        temp_map[col_idx] = matched
                        matches += 1
                if matches >= 2:
                    header_row_idx = idx
                    col_map = temp_map
                    break

        if header_row_idx == -1:
            header_row_idx = 0
            for col_idx, cell in enumerate(reader[0]):
                matched = _match_header_key(cell)
                if matched:
                    col_map[col_idx] = matched

        records = []
        for r in reader[header_row_idx + 1:]:
            if not any(c and c.strip() for c in r):
                continue
            item: Dict[str, Any] = {}
            for col_idx, key in col_map.items():
                if col_idx < len(r):
                    val = r[col_idx]
                    if key == "visit_date":
                        item[key] = _format_cell_date(val)
                    elif key == "visit_time":
                        item[key] = _format_cell_time(val)
                    elif key == "mobile":
                        item[key] = clean_mobile(str(val or ""))
                    else:
                        item[key] = str(val or "").strip()

            if item.get("candidate_name"):
                records.append(item)

        return records

    def import_visiting_register(self, filepath: str, repo=None) -> Dict[str, Any]:
        """
        Parses and commits Visiting Register records into local SQLite database.
        Returns a summary dictionary with counts.
        """
        active_repo = repo or getattr(self, "repo", repository)
        try:
            records = self.parse_visiting_register_file(filepath)
            if not records:
                return {
                    "success": False,
                    "total": 0,
                    "inserted": 0,
                    "skipped": 0,
                    "error": "No valid visitor entries with a candidate name were found in the file."
                }
            result = active_repo.bulk_import_visitors(records)
            return {
                "success": True,
                "total": len(records),
                "inserted": result["inserted"],
                "skipped": result["skipped"],
                "error": ""
            }
        except Exception as e:
            logger.error("Error during Visiting Register import: %s", e)
            return {
                "success": False,
                "total": 0,
                "inserted": 0,
                "skipped": 0,
                "error": str(e)
            }


file_importer = FileImporter()
